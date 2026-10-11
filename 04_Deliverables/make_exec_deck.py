# -*- coding: utf-8 -*-
"""Builds the 15-slide Executive Presentation from the full 30-slide deck.

The executive deck is a strict subset of PMARIS_Reliability_Deck.pptx (built by
make_deck.py), so every figure, chart and speaker note is identical between the
two. Slides are selected below, the rest are dropped, and the "n / 30" footers
are renumbered "n / 15".

Rebuild with:  python make_figures.py && python make_deck.py && python make_exec_deck.py
"""
import re
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
SRC = HERE / "PMARIS_Reliability_Deck.pptx"
OUT = HERE / "PMARIS_Executive_Presentation.pptx"

# 1-based slide numbers from the full deck, in presentation order
KEEP = [
    1,   # title
    2,   # the answer
    11,  # headline numbers
    12,  # concentration
    16,  # lever 1: strategy
    18,  # lever 2: spare parts
    19,  # levers compound (heatmap)
    20,  # lever 3: crew skill
    21,  # lever 4: PM discipline
    24,  # sizing (estimated, upper bound)
    25,  # escalation gap
    26,  # the proposed rule
    27,  # what does not matter
    29,  # recommendations
    30,  # decisions needed
]

prs = Presentation(str(SRC))
sld_list = prs.slides._sldIdLst
ids = list(sld_list)
keep_ids = [ids[i - 1] for i in KEEP]
for sid in ids:
    if sid not in keep_ids:
        prs.part.drop_rel(sid.rId)
        sld_list.remove(sid)
# reorder to KEEP order (already ascending, kept explicit for safety)
for sid in keep_ids:
    sld_list.remove(sid)
    sld_list.append(sid)

total = len(keep_ids)
pat = re.compile(r"^\d+ / 30$")
for n, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if pat.match(r.text.strip()):
                        r.text = f"{n} / {total}"
    if n == 1:
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.startswith("AIB902 · AGENTIC ANALYTICS"):
                            r.text = "EXECUTIVE PRESENTATION · AIB902 · GROUP 9230 TEAM 1"

prs.core_properties.title = "PMARIS Asset Reliability: Executive Presentation"
prs.save(str(OUT))
print("wrote", OUT, "slides:", total)
